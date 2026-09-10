"""Local, non-AI text extraction for the six target file types."""
import io
import os
import subprocess
import tempfile

import fitz  # PyMuPDF
import openpyxl
import xlrd
from docx import Document

MIN_PDF_TEXT_CHARS = 20  # below this, treat the PDF as scanned/no text layer


def extract_pdf_text(local_path: str) -> str:
    doc = fitz.open(local_path)
    text = "\n".join(page.get_text() for page in doc)
    doc.close()
    if len(text.strip()) >= MIN_PDF_TEXT_CHARS:
        return text
    return _ocr_pdf(local_path)


def _ocr_pdf(local_path: str) -> str:
    from pdf2image import convert_from_path
    import pytesseract

    pages = convert_from_path(local_path)
    return "\n".join(pytesseract.image_to_string(page) for page in pages)


def extract_docx_text(local_path: str) -> str:
    doc = Document(local_path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_doc_text(local_path: str) -> str:
    """Legacy binary .doc: convert to .docx with headless LibreOffice, then parse."""
    with tempfile.TemporaryDirectory() as tmpdir:
        subprocess.run(
            [
                "soffice", "--headless", "--norestore",
                "--convert-to", "docx", "--outdir", tmpdir, local_path,
            ],
            check=True, capture_output=True, timeout=120,
        )
        converted = os.path.join(
            tmpdir, os.path.splitext(os.path.basename(local_path))[0] + ".docx"
        )
        return extract_docx_text(converted)


def extract_xlsx_text(local_path: str) -> str:
    wb = openpyxl.load_workbook(local_path, data_only=True, read_only=True)
    parts = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_xls_text(local_path: str) -> str:
    wb = xlrd.open_workbook(local_path)
    parts = []
    for sheet in wb.sheets():
        for row_idx in range(sheet.nrows):
            cells = [str(c) for c in sheet.row_values(row_idx) if c not in ("", None)]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


_EXTRACTORS = {
    ".pdf": extract_pdf_text,
    ".docx": extract_docx_text,
    ".doc": extract_doc_text,
    ".xlsx": extract_xlsx_text,
    ".xlsm": extract_xlsx_text,
    ".xls": extract_xls_text,
}


def extract_text(local_path: str) -> str:
    ext = os.path.splitext(local_path)[1].lower()
    extractor = _EXTRACTORS.get(ext)
    if extractor is None:
        raise ValueError(f"Unsupported file extension: {ext}")
    return extractor(local_path)
